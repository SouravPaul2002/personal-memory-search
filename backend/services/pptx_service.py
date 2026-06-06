from pptx import Presentation


def extract_text_from_pptx(file_path):

    presentation = Presentation(file_path)

    text_parts = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    text_parts.append(text)

    return "\n".join(text_parts)